"""Tests for backend.m6_impact.jrc_damage: extraction from the JRC depth-damage
functions workbook into damage_curves.csv / asset_values.csv."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

from backend.m6_impact.jrc_damage import (
    JrcExtractionError,
    extract_asset_values,
    extract_damage_curves,
    write_all,
)
from backend.m6_impact.loss import load_loss_config

REPO_ROOT = Path(__file__).resolve().parents[2]
REAL_XLSX = REPO_ROOT / "data" / "copy_of_global_flood_depth-damage_functions__30102017.xlsx"

DAMAGE_ROWS = {
    "Residential buildings": [(0, 0.0), (1, 0.494), (6, 1.0)],
    "Commercial buildings": [(0, 0.0), (1, 0.5), (6, 1.0)],
    "Industrial buildings": [(0, 0.0), (1, 0.48), (6, 1.0)],
    "Transport": [(0, 0.0), (1, 0.57), (6, 1.0)],
    "Infrastructure - roads": [(0, 0.0), (1, 0.37), (6, 1.0)],
    "Agriculture": [(0, 0.0), (1, 0.37), (6, 1.0)],
}


def _build_workbook(tmp_path: Path) -> Path:
    """A minimal workbook with the same row/column layout as the real JRC file
    (docs/data_sources.md src_031/src_032), just fewer countries and depth points."""
    wb = Workbook()
    wb.remove(wb.active)

    df = wb.create_sheet("Damage functions")
    df.append([None] * 16)  # row 1: blank
    df.append(["Damage\nclass", "Flood depth,\n[m]", "Damage function", None, None, None, None, None, None,
               "Standard deviation"])  # row 2
    df.append([None, None, "EUROPE", "North AMERICA", "Centr&South\nAMERICA", "ASIA", "AFRICA", "OCEANIA",
                "GLOBAL"])  # row 3
    for label, points in DAMAGE_ROWS.items():
        for i, (depth, asia) in enumerate(points):
            df.append([label if i == 0 else None, depth, None, None, None, asia, None, None, None])

    def _building_sheet(name: str, india_total: float):
        ws = wb.create_sheet(name)
        ws.append([None, "Building based", None, None, "Land-use based", "Object based"])
        ws.append(["Country", "Max Damage Structure", "Max Damage Content", "Total", "Total", "Total"])
        ws.append([None, "(EUR/m2, 2010)", "(EUR/m2, 2010)", "(EUR/m2, 2010)", "(EUR/m2, 2010)", "(EUR/object, 2010)"])
        ws.append(["Elsewhere", 1.0, 1.0, 2.0, 1.0, 1000.0])
        ws.append(["India", india_total / 2, india_total / 2, india_total, india_total / 5, india_total * 100])

    _building_sheet("MaxDamage-Residential", 212.78)
    _building_sheet("MaxDamage-Commercial", 323.98)
    _building_sheet("MaxDamage-Industrial", 293.10)

    ag = wb.create_sheet("MaxDamage-Agriculture")
    ag.append([None, "Value Added/Hectare\n(avg 2008-2012)", "Area in km2 \n(avg 2008-2012)"])
    ag.append([None, "(EUR/ha, 2010)", "square km"])
    ag.append(["Elsewhere", 500.0, 1000.0])
    ag.append(["India", 1146.82, 1_797_058.0])

    def _gdp_sheet(name: str, region_ratio: float, region_gdp_ref: float, india_gdp: float):
        ws = wb.create_sheet(name)
        ws.append([None, None, None, None, None, None, None])
        ws.append(["Country", "GDP per capita (2010 US$)", None, None, None, None, "How to calculate"])
        ws.append(["Elsewhere", 1000.0, None, None, None, None, None, None, "Euro/m2", "GDP (2010 US$)"])
        ws.append(["Somewhere", 2000.0, None, None, None, None, "Maximum damage (average)", "Europe:", 25.0, 43000.0])
        ws.append(["India", india_gdp, None, None, None, None, None, "Asia:", region_ratio, region_gdp_ref])

    _gdp_sheet("MaxDamage-Infrastructure", region_ratio=4.0, region_gdp_ref=1913.0, india_gdp=1417.0736138018274)
    _gdp_sheet("MaxDamage-Transport", region_ratio=209.0, region_gdp_ref=2834.0, india_gdp=1417.0736138018274)

    path = tmp_path / "jrc_mini.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def mini_xlsx(tmp_path) -> Path:
    return _build_workbook(tmp_path)


@pytest.fixture
def config():
    return load_loss_config()


def test_extract_damage_curves_all_classes_and_asia_column(mini_xlsx):
    rows = extract_damage_curves(mini_xlsx)
    classes = {r["asset_class"] for r in rows}
    assert classes == {"residential", "commercial", "industrial", "transport", "infrastructure_roads", "agriculture"}
    residential = sorted((r["depth_m"], r["damage_fraction"]) for r in rows if r["asset_class"] == "residential")
    assert residential == [(0.0, 0.0), (1.0, 0.494), (6.0, 1.0)]
    assert all(r["source"] == "src_031" for r in rows)


def test_extract_damage_curves_missing_class_raises(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    df = wb.create_sheet("Damage functions")
    df.append([None] * 10)
    df.append(["Damage\nclass", "Flood depth,\n[m]", "Damage function", None, None, None, None, None, None,
               "Standard deviation"])
    df.append([None, None, "EUROPE", "North AMERICA", "Centr&South\nAMERICA", "ASIA"])
    df.append(["Residential buildings", 0, 0, 0, 0, 0])
    path = tmp_path / "incomplete.xlsx"
    wb.save(path)
    with pytest.raises(JrcExtractionError):
        extract_damage_curves(path)


def test_extract_asset_values_india_building_based_total(mini_xlsx, config):
    rows = {r["asset_class"]: r for r in extract_asset_values(mini_xlsx, config)}
    assert rows["residential"]["value_eur2010"] == pytest.approx(212.78)
    assert rows["commercial"]["value_eur2010"] == pytest.approx(323.98)
    assert rows["industrial"]["value_eur2010"] == pytest.approx(293.10)
    assert "MaxDamage-Residential" in rows["residential"]["jrc_cell"]


def test_extract_asset_values_agriculture_converts_ha_to_m2(mini_xlsx, config):
    rows = {r["asset_class"]: r for r in extract_asset_values(mini_xlsx, config)}
    assert rows["agriculture"]["value_eur2010"] == pytest.approx(1146.82 / 10_000.0)


def test_extract_asset_values_infrastructure_and_transport_use_gdp_ratio(mini_xlsx, config):
    rows = {r["asset_class"]: r for r in extract_asset_values(mini_xlsx, config)}
    india_gdp = 1417.0736138018274
    assert rows["infrastructure_roads"]["value_eur2010"] == pytest.approx(india_gdp * (4.0 / 1913.0))
    assert rows["transport"]["value_eur2010"] == pytest.approx(india_gdp * (209.0 / 2834.0))


def test_extract_asset_values_placeholder_fx_leaves_inr_blank(mini_xlsx, config):
    assert config.eur_to_inr_2010.status == "placeholder"  # config/impact.yaml ships unsourced
    rows = {r["asset_class"]: r for r in extract_asset_values(mini_xlsx, config)}
    assert all(r["value_inr_per_unit"] == "" for r in rows.values())
    assert all(r["status"] == "placeholder" for r in rows.values())


def test_write_all_writes_both_csvs_with_contract_columns(mini_xlsx, tmp_path, config):
    out = write_all("testsite", mini_xlsx, data_dir=tmp_path, config=config)
    with open(out["damage_curves"], newline="", encoding="utf-8") as f:
        header = next(csv.reader(f))
    assert header == ["asset_class", "depth_m", "damage_fraction", "source"]
    with open(out["asset_values"], newline="", encoding="utf-8") as f:
        header = next(csv.reader(f))
    assert header == ["asset_class", "value_inr_per_unit", "unit", "source", "status", "value_eur2010", "jrc_cell"]


@pytest.mark.skipif(not REAL_XLSX.is_file(), reason="real JRC workbook is gitignored (data/), not present here")
def test_real_workbook_india_residential_matches_cited_cell():
    """Confirms the extractor reads the real, published JRC figures, not a
    stand-in (docs/data_sources.md src_032)."""
    rows = {r["asset_class"]: r for r in extract_asset_values(REAL_XLSX, load_loss_config())}
    assert rows["residential"]["value_eur2010"] == pytest.approx(212.7798639329916)
    assert rows["residential"]["jrc_cell"] == "'MaxDamage-Residential'!D90"
